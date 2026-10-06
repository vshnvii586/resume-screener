import unittest
from fastapi.testclient import TestClient
import docx
import io

from backend.app import app

class TestExtractionAPI(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        
    def test_docx_resume_extraction(self):
        # 1. Construct an actual DOCX resume.
        doc = docx.Document()
        doc.add_paragraph("Alex Sharma")
        doc.add_paragraph("Software Engineer")
        doc.add_paragraph("Experience: 5 years in Python")
        
        file_bytes = io.BytesIO()
        doc.save(file_bytes)
        file_bytes.seek(0)
        
        # 2. Send it through the API endpoint.
        response = self.client.post(
            "/api/resumes/ai-parse",
            files={"file": ("resume.docx", file_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )
        
        # 3. Verify the endpoint succeeds.
        self.assertEqual(response.status_code, 200)
        
        data = response.json()
        
        # 4. Verify structured resume data exists.
        self.assertIn("profile", data)
        self.assertIn("name", data["profile"])
        
        # 5. Verify the extracted text is non-empty.
        self.assertIn("raw_text", data["profile"])
        self.assertTrue(len(data["profile"]["raw_text"]) > 0)
        self.assertIn("Alex Sharma", data["profile"]["raw_text"])
        
        # 6. Verify the frontend-compatible response structure.
        self.assertEqual(data["filename"], "resume.docx")
        
