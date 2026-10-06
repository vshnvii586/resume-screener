import unittest
import json
import io
import docx
from backend.extractor import process_extraction

class TestExtractor(unittest.TestCase):
    def setUp(self):
        self.resume_text = """
Jordan Lee
your.name@gmail.com
(206) 555-0184

SUMMARY
Operations Analyst with 6+ years of experience optimizing processes.

EDUCATION
Bachelor of Business Administration
University of Washington

PROFESSIONAL EXPERIENCE
Senior Operations Analyst
NovaBridge Solutions - Seattle, WA
Jan 2019 - Present
- Led implementation of a new internal workflow system

SKILLS
Project Management, Stakeholder management, Process improvement
        """
        
        self.jd_text = """
Job Title: Business Operations Analyst

Requirements:
- Bachelor's degree in Business Administration
- 3+ years of experience in business operations
- Strong process improvement skills
- Deep experience in data analysis
- Microsoft Excel

Preferred Qualifications:
- Experience with SQL
- Familiarity with Power BI
        """

    def test_docx_resume(self):
        # Create a valid DOCX file in memory
        doc = docx.Document()
        for line in self.resume_text.strip().split('\n'):
            doc.add_paragraph(line)
            
        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        doc_bytes = doc_stream.getvalue()
        
        result = process_extraction(doc_bytes, "docx", self.jd_text)
        
        resume = result["resume"]
        jd = result["job_description"]
        
        # Verify resume
        self.assertEqual(resume["name"], "Jordan Lee")
        self.assertEqual(resume["email"], "your.name@gmail.com")
        self.assertEqual(resume["phone"], "(206) 555-0184")
        self.assertEqual(len(resume["experience"]), 1)
        self.assertEqual(resume["experience"][0]["title"], "Senior Operations Analyst")
        self.assertEqual(resume["experience"][0]["company"], "NovaBridge Solutions")
        self.assertTrue(resume["education"])
        self.assertIn("Process improvement", resume["skills"]["technical"])
        self.assertTrue(resume["raw_text"])
        
        # Verify JD
        self.assertEqual(jd["job_title"], "Job Title: Business Operations Analyst")
        self.assertIn("process improvement", jd["required_skills"])
        self.assertIn("SQL", jd["preferred_skills"])
        self.assertEqual(jd["required_experience"], "3+ years")

    def test_pdf_resume(self):
        # We simulate PDF extraction by just mocking the text extractor since generating a valid PDF from scratch is non-trivial without extra deps
        import backend.extractor.resume_extractor as re_ext
        original_extract = re_ext.extract_text_from_file_bytes
        re_ext.extract_text_from_file_bytes = lambda b, e: self.resume_text
        
        try:
            result = process_extraction(b"fake_pdf_bytes", "pdf", self.jd_text)
            self.assertEqual(result["resume"]["name"], "Jordan Lee")
            self.assertTrue(result["resume"]["raw_text"])
        finally:
            re_ext.extract_text_from_file_bytes = original_extract

    def test_missing_jd(self):
        # Missing JD should return empty jd dict
        doc = docx.Document()
        doc.add_paragraph(self.resume_text)
        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        
        result = process_extraction(doc_stream.getvalue(), "docx", "")
        self.assertEqual(result["job_description"], {})
        self.assertEqual(result["resume"]["name"], "Jordan Lee")

    def test_unsupported_resume(self):
        # Unsupported resume format (e.g. txt)
        doc_bytes = b"Just some text"
        # Since extract_text_from_file_bytes only supports pdf and docx, it will return empty string for 'txt'
        result = process_extraction(doc_bytes, "txt", self.jd_text)
        self.assertEqual(result["resume"]["raw_text"], "")
        self.assertEqual(result["resume"]["name"], "")

    def test_resume_missing_sections(self):
        # Resume with just a name and email, no sections
        minimal_resume = "Alex Smith\nalex@example.com\n"
        doc = docx.Document()
        doc.add_paragraph(minimal_resume)
        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        
        result = process_extraction(doc_stream.getvalue(), "docx", "")
        self.assertEqual(result["resume"]["name"], "Alex Smith")
        self.assertEqual(result["resume"]["email"], "alex@example.com")
        self.assertEqual(result["resume"]["experience"], [])
        self.assertEqual(result["resume"]["education"], [])

if __name__ == "__main__":
    unittest.main()
